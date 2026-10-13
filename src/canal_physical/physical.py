"""The experimental method rebuilt around the physical scale.

At 0.13 um/px a canaliculus of 0.2 to 0.4 um is about 1.5 to 3 px across, and
about 2.5 to 4 px after the optical blur. The masks of the current method have a
median distance transform width of 6.0 px (0.78 um) and the first ridge variants
7.4 px, so both are roughly twice that. The scales of those first variants came
from a half width measured on the already inflated mask, which is circular
reasoning: the mask set the scale and the scale set the mask.

Here every scale is set in micrometres from the pixel size, or chosen from the
image itself, and never from the old mask and never from a published value.

**Pixel size 0.13 um/px.** Pre-validation: no number here has been checked
against a manual count. A diagnosis and a better founded experiment, not a claim
that anything is better.

The flattening, the lacuna buffer, the vascular handling, the bridging rule, the
graph cleanup, the ownership and the ring radii stay exactly as the current
method has them, so one idea changes at a time. The top-hat radius of the
flattening is 5 px, which is 0.65 um and wider than a thread; it is reported with
every run and left alone here, and remains an open question.

Three variants are kept, reported in this order: P5 is the ridge method with gap
bridging off, A is the current method as it stands, and P2 is P5 with gap
bridging on, so P2 and P5 differ by that one switch and nothing else. P5 leads
because it is the preferred variant by eye, which is a visual preference and not
a validated result. The output is one folder, results/ridge/: the comparison
tables and, per image, the three variants side by side as one png.
"""
from __future__ import annotations

import csv
import json
import sys
from dataclasses import replace
from pathlib import Path

_ROOT = Path(__file__).resolve().parents[2]
for _extra in (str(_ROOT), str(_ROOT / "src")):
    if _extra not in sys.path:
        sys.path.insert(0, _extra)

import numpy as np  # noqa: E402
from skimage import morphology, segmentation  # noqa: E402

import canaliculi  # noqa: E402
import config
import lacunae
from canal_physical import metrics, params as params_mod, preprocess, ridge
from canal_physical import skeleton
from canal_physical import widthprofile

PIXEL_SIZE_UM = config.PIXEL_SIZE_UM_CANAL_PHYSICAL

# (key, what it is, family, bridging, sigmas in um or None, low cut share)
# The ridge filter always reads the flattened channel. P2 and P5 carry the same
# scales and the same cut and differ only in the bridging switch.
#
# P5 is listed first because it is the preferred variant by eye. That is a visual
# preference and nothing more: it is NOT validated, and no variant is called
# better than another by any number here. The order decides how the tables, the
# report and the figure read; it decides nothing that is computed, since every
# variant is looked up by key and every change is still measured against A.
VARIANTS = [
    ("P5", "ridge, sigmas 0.17, 0.21, 0.26 um, gap bridging off", "ridge", False,
     (0.17, 0.21, 0.26), 0.75),
    ("A", "current method, as is", "current", True, None, None),
    ("P2", "ridge, sigmas 0.17, 0.21, 0.26 um, gap bridging on", "ridge", True,
     (0.17, 0.21, 0.26), 0.75),
]
VARIANT_KEYS = [v[0] for v in VARIANTS]
VARIANT_LABELS = {v[0]: v[1] for v in VARIANTS}

def build(key: str, lac: dict, channel: np.ndarray, label: str, params) -> dict:
    """One variant's detection. Everything except the thresholded quantity, the
    scales and the bridging switch is the current method's own code."""
    _k, what, family, bridging, sigmas_um, low_fraction = next(v for v in VARIANTS if v[0] == key)
    lacuna_mask, lacuna_id_map = canaliculi.build_lacuna_maps(lac["labels"], lac["kept"])
    flagged = preprocess.vascular_mask(channel)
    flattened = preprocess.flatten(channel)
    lacuna_parameters = {**lacunae.parameters(), "computed_threshold_t_hi": lac["t_hi"]}

    if family == "current":
        candidate, cut = canaliculi.network_candidate_mask(flattened, lacuna_mask, flagged)
        thresholded = candidate
        skel = morphology.skeletonize(candidate)
        bridges = []
        if bridging:
            bridges = canaliculi.find_bridges(skel, flattened, cut, lacuna_mask | flagged)
            if bridges:
                candidate = canaliculi.apply_bridges(candidate, bridges)
                skel = morphology.skeletonize(candidate)
        network = {"variant": key, "what": what, "thresholded_quantity": "flattened intensity",
                   "gap_bridging": bridging, "high_cut": float(cut),
                   "low_cut": float(cut) * canaliculi.HYSTERESIS_LOW_FRACTION,
                   "ridge_source_image": None}
        source_image = flattened
    else:
        source_image = flattened
        variant_params = replace(params, hysteresis_low_fraction=low_fraction)
        sigmas_px = tuple(params.px(s) for s in sigmas_um)
        response = ridge.ridge_response(source_image, variant_params, sigmas_px)
        candidate, cut = skeleton.ridge_network_mask(response, lacuna_mask, flagged, variant_params)
        thresholded = candidate
        skel = morphology.skeletonize(candidate)
        bridges = []
        if bridging:
            bridges = canaliculi.find_bridges(skel, response, cut, lacuna_mask | flagged)
            if bridges:
                candidate = canaliculi.apply_bridges(candidate, bridges)
                skel = morphology.skeletonize(candidate)
        network = {
            "variant": key, "what": what,
            "thresholded_quantity": "multiscale ridge response (skimage.filters.sato, bright ridges)",
            "gap_bridging": bridging, "sigmas_um": list(sigmas_um),
            "sigmas_px": [round(s, 4) for s in sigmas_px],
            "hysteresis_low_fraction": low_fraction, "high_cut": float(cut),
            "low_cut": float(cut) * low_fraction,
            "ridge_source_image": "flattened channel",
        }

    dist_to_lacuna, nearest_id = canaliculi.nearest_lacuna_map(lacuna_id_map)
    own = canaliculi.build_ownership(lac["kept"], skel, dist_to_lacuna, nearest_id)
    detection = metrics.detection_dict(label, lacuna_id_map, candidate, candidate & ~thresholded, skel,
                                       flagged, own, len(bridges), lacuna_parameters, network)
    detection["owner_map"] = own["owner_map"]
    detection["bridges"] = bridges
    detection["flattened"] = flattened
    detection["source_image"] = source_image
    detection["lacuna_mask"] = lacuna_mask
    return detection


def reference(image_path, lac: dict, channel: np.ndarray) -> dict:
    """Variant A straight from the pipeline, so the reference is the pipeline and
    not a copy of it."""
    import quantification

    detected = canaliculi.analyse_network(image_path, lac, channel)
    detection = quantification.detection_in_memory(detected)
    detection["owner_map"] = detected["owner_map"]
    detection["bridges"] = detected["bridges"]
    detection["flattened"] = canaliculi.preprocess_channel(channel)
    detection["source_image"] = detection["flattened"]
    detection["lacuna_mask"] = detection["lacuna_id_map"] > 0
    return detection


# The measures
# Everything the pipeline already measures is measured by src/quantification.py,
# the same code for every variant. The width on the image, the contrast and the
# smoothed straight runs are added here.

def usable_for_width(detection: dict, params) -> np.ndarray:
    """The skeleton pixels a width profile may be taken at: on the thresholded
    mask, not drawn by bridging, outside the lacuna buffer, outside the vascular
    mask with its dilation, and not a junction."""
    skel = detection["skeleton"]
    buffer_px = params.px_int(params.lacuna_buffer_um)
    buffered = morphology.dilation(detection["lacuna_mask"], morphology.disk(buffer_px))
    return (skel & detection["thresholded"] & ~detection["bridged"] & ~buffered
            & ~detection["vascular"] & ~widthprofile.junction_mask(skel))


def smoothed_straight_run_fraction(skel: np.ndarray, window: int = 5,
                                   min_run: int = metrics.STRAIGHT_RUN_MIN_PX,
                                   tolerance_px: float = 0.5) -> float | None:
    """The share of skeleton path points that lie in an axis-aligned run after the
    path is smoothed.

    A one pixel line drawn on a pixel grid has to step, so a straight run counter
    on the raw pixels partly measures the grid and not the structure. Here each
    branch of the skeleton is taken as an ordered path (skan, the same library the
    pipeline uses to build its graph), the path coordinates are smoothed with a
    moving average of `window` points, and a point counts when it lies in a stretch
    of at least `min_run` consecutive smoothed points whose rows, or whose columns,
    span less than `tolerance_px`. A staircase that is really a diagonal smooths
    into a diagonal and is not counted; a genuinely straight run stays straight."""
    from skan import Skeleton

    if not skel.any():
        return None
    try:
        paths = Skeleton(skel)
    except Exception:  # noqa: BLE001  a skeleton with no branch at all
        return None
    total = marked = 0
    box = np.ones(window) / window
    for index in range(paths.n_paths):
        coords = paths.path_coordinates(index)
        n = len(coords)
        total += n
        if n < min_run:
            continue
        smooth = np.empty_like(coords, dtype=float)
        for axis in (0, 1):
            padded = np.pad(coords[:, axis].astype(float), (window // 2, window // 2), mode="edge")
            smooth[:, axis] = np.convolve(padded, box, mode="valid")[:n]
        in_run = np.zeros(n, dtype=bool)
        for start in range(0, n - min_run + 1):
            stop = start + min_run
            piece = smooth[start:stop]
            if (np.ptp(piece[:, 0]) < tolerance_px) or (np.ptp(piece[:, 1]) < tolerance_px):
                in_run[start:stop] = True
        marked += int(in_run.sum())
    return round(marked / total, 6) if total else None


def measure_variant(key: str, detection: dict, image_path, params, legacy_cut: float,
                    precision: int = 4) -> dict:
    """Every measure of one variant on one image."""
    measured = metrics.measure(detection, image_path, params)
    skel = detection["skeleton"]

    # Width on the image.
    usable = usable_for_width(detection, params)
    profile = widthprofile.measure(detection["flattened"], skel, usable)
    width = widthprofile.summarise(profile["width_px"], profile["ok"], profile["n_candidates"],
                                   params.pixel_size_um, precision)

    # Contrast against the noise of this variant's own background.
    exclude = detection["lacuna_mask"] | detection["vascular"] | detection["thresholded"]
    noise = widthprofile.noise_level(detection["flattened"], exclude)
    snr = widthprofile.contrast(detection["flattened"], profile["rows"], profile["cols"],
                                profile["normal_r"], profile["normal_c"], noise)
    finite = snr[np.isfinite(snr)]
    contrast = {
        "noise_level": round(noise, 8),
        "contrast_snr_median": round(float(np.median(finite)), precision) if finite.size else None,
        "contrast_snr_p10": round(float(np.percentile(finite, 10)), precision) if finite.size else None,
        "low_contrast_fraction": (round(float((finite < widthprofile.LOW_CONTRAST_SNR).mean()), precision)
                                  if finite.size else None),
        "contrast_pixels_used": int(finite.size),
    }

    # Artifact measures: the straight runs raw and smoothed, the rectangles, and
    # the legacy support measure kept only for continuity.
    artifacts = metrics.artifact_measures(detection, detection["flattened"], legacy_cut)
    artifacts["legacy_unsupported_fraction"] = artifacts.pop("unsupported_fraction")
    artifacts["legacy_unsupported_px"] = artifacts.pop("unsupported_px")
    artifacts["straight_run_fraction_smoothed"] = smoothed_straight_run_fraction(skel)

    return {"variant": key, "measured": measured, "width": width, "contrast": contrast,
            "artifacts": artifacts, "profile": profile, "usable": usable}


MEASURE_COLUMNS = [
    "image_label", "variant", "what", "pixel_size_label",
    # The width, on the image and on the mask, side by side.
    "width_fwhm_median_px", "width_fwhm_median_um", "width_fwhm_p10_px", "width_fwhm_p10_um",
    "width_fwhm_p90_px", "width_fwhm_p90_um", "width_fwhm_pixels_used", "width_fwhm_failure_fraction",
    "width_dt_median_px", "width_dt_median_um", "width_dt_p10_px", "width_dt_p90_px",
    # Support, neutral and legacy.
    "contrast_snr_median", "contrast_snr_p10", "low_contrast_fraction", "contrast_pixels_used",
    "noise_level", "legacy_unsupported_fraction",
    # Shape of the traced network.
    "straight_run_fraction", "straight_run_fraction_smoothed", "rectangle_count",
    "skeleton_px", "bridged_fraction", "n_bridges",
    # Everything the pipeline measures.
    "field_length_density_per_px", "field_length_density_um_per_um2",
    "roots_per_cell", "ring_length_r30_px_per_cell", "ring_length_r30_um_per_cell",
    "ring_length_r60_px_per_cell", "ring_length_r60_um_per_cell",
    "junction_count", "junction_density_per_um2", "thread_end_fraction",
    "connected_to_lacuna_fraction", "median_edge_length_px", "median_edge_length_um",
    "mean_edge_length_px", "mean_edge_length_um",
    "lacuna_count", "interior_lacuna_count",
    # What the variant was.
    "sigmas_um", "sigmas_px", "hysteresis_low_fraction", "gap_bridging", "ridge_source_image",
]

SUMMARY_MEASURES = [
    "width_fwhm_median_um", "width_dt_median_um", "low_contrast_fraction", "contrast_snr_median",
    "straight_run_fraction", "straight_run_fraction_smoothed", "rectangle_count",
    "connected_to_lacuna_fraction", "junction_count", "thread_end_fraction",
    "field_length_density_um_per_um2", "roots_per_cell", "skeleton_px", "n_bridges",
    "legacy_unsupported_fraction",
]


def measure_row(label: str, result: dict, detection: dict) -> list:
    """One row of the comparison table, in MEASURE_COLUMNS order."""
    comparison = result["measured"]["comparison"]
    network = detection["canaliculi_detection"]["parameters"]
    row = {
        "image_label": label,
        "variant": result["variant"],
        "what": VARIANT_LABELS[result["variant"]],
        "pixel_size_label": params_mod.PIXEL_LABEL,
        "width_dt_median_px": comparison.get("width_median_px"),
        "width_dt_median_um": comparison.get("width_median_um"),
        "width_dt_p10_px": comparison.get("width_p10_px"),
        "width_dt_p90_px": comparison.get("width_p90_px"),
        "sigmas_um": ", ".join(str(s) for s in network.get("sigmas_um", [])) or None,
        "sigmas_px": ", ".join(str(s) for s in network.get("sigmas_px", [])) or None,
        "hysteresis_low_fraction": network.get("hysteresis_low_fraction"),
        "gap_bridging": network.get("gap_bridging"),
        "ridge_source_image": network.get("ridge_source_image"),
        **result["width"], **result["contrast"], **result["artifacts"], **comparison,
    }
    return [row.get(c) for c in MEASURE_COLUMNS]


# Running it

OUT_DIR = config.RESULTS_DIR / "ridge"
# The comparison holds the three variants side by side, so it is drawn at a
# smaller upscale than a single panel would be to keep the file in budget.
COMPARE_UPSCALE = 2


def compare_with_results(label: str, measured: dict) -> list:
    """Variant A against results/<label>/5_quantification/, at tolerance 0."""
    import quantification

    committed = json.loads(quantification.quant_path(config.RESULTS_DIR, label, ".json")
                           .read_text(encoding="utf-8"))
    diffs = []

    def walk(a, b, where):
        if isinstance(a, dict):
            for key, value in a.items():
                walk(value, (b or {}).get(key), f"{where}.{key}" if where else key)
        elif isinstance(a, list):
            for i, value in enumerate(a):
                walk(value, b[i] if isinstance(b, list) and i < len(b) else None, f"{where}[{i}]")
        else:
            if isinstance(a, bool) or isinstance(b, bool):
                same = a == b
            elif isinstance(a, (int, float)) and isinstance(b, (int, float)):
                same = a == b or (a != a and b != b)
            else:
                same = a == b
            if not same:
                diffs.append((where, a, b))

    for key, value in (("lacuna_count", measured["lacuna_count"]),
                       ("interior_lacuna_count", measured["interior_lacuna_count"]),
                       ("n_bridges", measured["n_bridges"]),
                       ("summary", measured["summary"]), ("field", measured["field"]),
                       ("width", measured["width"]), ("lacunae", measured["rows"])):
        walk(committed.get(key), value, key)
    return diffs


def verification_matches_pipeline(label: str, path: Path) -> tuple:
    from skimage.io import imread

    committed = lacunae.result_path(config.RESULTS_DIR, label, config.SECTION_CANALICULI,
                                    config.SUFFIX_CANALICULI_VERIFICATION)
    if not committed.is_file():
        return False, f"{committed} is missing"
    a, b = np.asarray(imread(path)), np.asarray(imread(committed))
    if a.shape != b.shape or not np.array_equal(a, b):
        return False, "differs from the committed verification image"
    return True, "pixel identical"


def write_figures(label: str, display, detections: dict, lac: dict, params,
                  out_dir: Path) -> dict:
    """The one figure of one image: the three variants side by side, as a png.
    Returns what the checks need.

    Variant A's pipeline style verification picture is drawn too, compared with
    the committed one in results/<label>/2_canaliculi/, and then deleted. It is a
    byte for byte duplicate of a file the repository already holds, so keeping it
    would only cost space, but drawing it is the check that variant A really is
    the pipeline."""
    import tempfile

    folder = out_dir / label
    colors = canaliculi.lacuna_colors(len(lac["kept"]))

    panels = []
    for key in VARIANT_KEYS:
        panel = draw_skeleton(display, detections[key]["skeleton"], detections[key]["owner_map"],
                              detections[key]["lacuna_id_map"], colors, COMPARE_UPSCALE)
        panels.append(add_scale_bar(panel, params.pixel_size_um, COMPARE_UPSCALE))
    gap = np.zeros((panels[0].shape[0], 8, 3), dtype=panels[0].dtype)
    both = panels[0]
    for panel in panels[1:]:
        both = np.concatenate([both, gap, panel], axis=1)
    reserved = list(colors.values()) + [UNOWNED_COLOUR, (255, 255, 255)]
    save_png(both, folder / f"{label}_compare.png", reserved)

    temporary = Path(tempfile.gettempdir()) / f"{label}_verification_A_check.png"
    pipeline_style(lac, display, detections["A"], temporary)
    checks = {"verification": verification_matches_pipeline(label, temporary)}
    temporary.unlink(missing_ok=True)
    return checks


def fixed_workbook(wb, path: Path) -> None:
    """Save a workbook whose bytes depend only on its contents, and check it
    opens afterwards."""
    import datetime
    import re
    import shutil
    import zipfile

    from openpyxl import load_workbook

    stamp = datetime.datetime(1980, 1, 1)
    wb.properties.created = stamp
    wb.properties.modified = stamp
    wb.properties.creator = ""
    wb.properties.lastModifiedBy = ""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    fixed = b"1980-01-01T00:00:00Z"
    tmp = path.with_name(path.name + ".tmp")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in sorted(src.infolist(), key=lambda i: i.filename):
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*", rb"\g<1>" + fixed, data)
            entry = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = info.external_attr
            dst.writestr(entry, data)
    shutil.move(str(tmp), str(path))
    load_workbook(path)  # it must open, or this raises here rather than for the reader


def write_tables(rows: list, out_dir: Path, params) -> dict:
    out_dir.mkdir(parents=True, exist_ok=True)
    with open(out_dir / "comparison_all_images.csv", "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(MEASURE_COLUMNS)
        writer.writerows(rows)

    index = {c: i for i, c in enumerate(MEASURE_COLUMNS)}
    means = {}
    for key in VARIANT_KEYS:
        taken = [r for r in rows if r[index["variant"]] == key]
        means[key] = {}
        for name in SUMMARY_MEASURES:
            values = [r[index[name]] for r in taken
                      if isinstance(r[index[name]], (int, float)) and r[index[name]] is not None]
            means[key][name] = float(np.mean(values)) if values else None

    summary_rows = []
    for key in VARIANT_KEYS:
        for name in SUMMARY_MEASURES:
            value, base = means[key][name], means["A"][name]
            change = (round(100.0 * (value - base) / base, 2)
                      if value is not None and base not in (None, 0) else None)
            summary_rows.append([key, VARIANT_LABELS[key], name,
                                 None if value is None else round(value, 6),
                                 None if base is None else round(base, 6), change])
    with open(out_dir / "comparison_summary.csv", "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["variant", "what", "measure", "mean_over_images", "variant_A_mean",
                         "percent_change_from_A"])
        writer.writerows(summary_rows)

    from openpyxl import Workbook

    wb = Workbook()
    sheet = wb.active
    sheet.title = "per_image"
    sheet.append(MEASURE_COLUMNS)
    for row in rows:
        sheet.append(row)
    s = wb.create_sheet("summary")
    s.append(["variant", "what", "measure", "mean_over_images", "variant_A_mean",
              "percent_change_from_A"])
    for row in summary_rows:
        s.append(row)
    notes = wb.create_sheet("notes")
    for line in notes_lines(params):
        notes.append([line])
    fixed_workbook(wb, out_dir / "comparison_all_images.xlsx")
    return means


def notes_lines(params) -> list:
    return [
        params_mod.PIXEL_LABEL + ".",
        "Pre-validation: no number here has been checked against a manual count.",
        "A diagnosis and a better founded experiment. No variant is called better than another.",
        "width_fwhm_*: the full width at half maximum of the intensity profile across the thread, taken",
        "  on the flattened image along the normal at each usable skeleton pixel. It INCLUDES the optical",
        "  blur, so it is an apparent width in the image and an upper bound on the diameter, not the",
        "  diameter. No point spread function is computed or subtracted: the numerical aperture, the",
        "  emission wavelength and the pinhole are unknown.",
        "width_dt_*: the older measure, twice the distance transform of the mask, kept beside it. It can",
        "  only return the width of the mask it is given.",
        "contrast_snr: the intensity at a skeleton pixel minus the median intensity 4 to 6 px away along",
        "  the normal on both sides, over the noise, where the noise is 1.4826 times the median absolute",
        "  deviation of the flattened image outside the lacunae, the vascular regions and that variant's",
        "  own mask. low_contrast_fraction is the share below 2.",
        "legacy_unsupported_fraction: the earlier measure, which used the CURRENT method's own threshold",
        "  and so favoured the current method. Kept only so the earlier tables can be followed.",
        "straight_run_fraction: share of skeleton pixels in an axis-aligned run of at least 8 px. A thin",
        "  skeleton on a pixel grid contains straight runs by its nature, and the current method already",
        "  scores about 0.22, so only differences between variants mean anything here.",
        "straight_run_fraction_smoothed: the same after each branch path is smoothed with a 5 point",
        "  moving average, which removes most of the pixel grid effect.",
    ]


def main() -> None:
    import argparse

    parser = argparse.ArgumentParser(
        description="The canalicular experiment rebuilt around the physical scale "
                    "(pre-validation, pixel size 0.13 um/px). Not a claim of improvement.")
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--image", type=Path, help="One .tif image.")
    group.add_argument("--dir", type=Path, help="A folder of .tif images.")
    parser.add_argument("-o", "--out", type=Path, default=OUT_DIR, help="Output folder.")
    args = parser.parse_args()

    params = params_mod.DEFAULT
    # Every micrometre parameter must convert back to exactly the pixel value the
    # imported pipeline code uses, or the variants are not comparable. This
    # raises rather than reporting a number built on a mismatch.
    checked = skeleton.check_parameters_match_current(params)
    print(f"Pixel size {params.pixel_size_um} um/px. {params_mod.PIXEL_LABEL}.")
    print(f"{len(checked)} parameters match the current method's pixel values exactly.")
    print(f"Field 1024 px = {1024 * params.pixel_size_um:g} um.")
    for key in VARIANT_KEYS:
        print(f"  {key:3s} {VARIANT_LABELS[key]}")
    print()

    image_paths = {lacunae.image_label(p): p for p in
                   ([args.image] if args.image else sorted(args.dir.glob("*.tif")))}
    rows, failures = [], []
    for label, path in image_paths.items():
        lac = lacunae.analyse_image(path)
        display, channel = lacunae.load_channel(path)

        detections, results = {}, {}
        for key in VARIANT_KEYS:
            detections[key] = (reference(path, lac, channel) if key == "A"
                               else build(key, lac, channel, label, params))
        _m, legacy_cut = canaliculi.total_signal_mask(detections["A"]["flattened"])
        for key in VARIANT_KEYS:
            results[key] = measure_variant(key, detections[key], path, params, legacy_cut)
            rows.append(measure_row(label, results[key], detections[key]))

        diffs = compare_with_results(label, results["A"]["measured"])
        checks = write_figures(label, display, detections, lac, params, args.out)
        checks["variant_A_against_results"] = {"ok": not diffs, "differences": len(diffs)}
        checks["variant_A_verification"] = {"ok": checks["verification"][0],
                                            "detail": checks["verification"][1]}
        checks.pop("verification")
        if diffs:
            failures.append((label, "variant A differs from results/", len(diffs)))
        if not checks["variant_A_verification"]["ok"]:
            failures.append((label, "variant A verification image differs",
                             checks["variant_A_verification"]["detail"]))
        print(f"{label}: A against results/ {'PASS' if not diffs else f'FAIL {len(diffs)}'}; "
              f"A verification {checks['variant_A_verification']['detail']}")
        for key in VARIANT_KEYS:
            w = results[key]["width"]
            a = results[key]["artifacts"]
            c = results[key]["contrast"]
            print(f"   {key:3s} fwhm {str(w['width_fwhm_median_um']):>7s} um  dt "
                  f"{str(results[key]['measured']['comparison'].get('width_median_um')):>7s} um  "
                  f"low contrast {str(c['low_contrast_fraction']):>7s}  straight "
                  f"{str(a['straight_run_fraction']):>7s} / {str(a['straight_run_fraction_smoothed']):>7s}")

    means = write_tables(rows, args.out, params)
    write_readme(args.out, params, image_paths, means)

    print(f"\ntable -> {args.out / 'comparison_all_images.xlsx'} and .csv, "
          f"summary -> {args.out / 'comparison_summary.csv'}")
    if failures:
        print("FAIL:")
        for f in failures:
            print("  ", *f)
        raise SystemExit(1)
    print("PASS: variant A equals results/ at tolerance 0 and its pipeline style verification image "
          "is pixel identical")




def write_readme(out_dir: Path, params, image_paths: dict, means: dict) -> None:
    import platform
    from importlib import metadata

    versions = {"python": platform.python_version()}
    for package in ("numpy", "scipy", "scikit-image", "networkx", "skan", "matplotlib", "openpyxl",
                    "pillow"):
        try:
            versions[package] = metadata.version(package)
        except metadata.PackageNotFoundError:
            versions[package] = None
    prov = lacunae.provenance()
    lines = [
        "# The canalicular experiment on the physical scale",
        "",
        f"**{params_mod.PIXEL_LABEL}.** Pre-validation: no number here has been checked against a",
        "manual count. **No variant is called better than another**, and nothing was tuned to improve",
        "any result or to approach any published value.",
        "",
        "## Why this exists",
        "",
        f"At {params.pixel_size_um} um/px a canaliculus of 0.2 to 0.4 um is about 1.5 to 3 px across, and about",
        "2.5 to 4 px once the optics have blurred it. The masks of the current method have a median",
        "distance transform width of 6.0 px (0.78 um) and the first ridge variants 7.4 px, so both are",
        "roughly twice that. The scales of those first variants came from a half width measured on the",
        "already inflated mask, which is circular: the mask set the scale and the scale set the mask.",
        "",
        "Here every scale is set in micrometres from the pixel size, and the width is measured on the",
        "image rather than on the mask.",
        "",
        "## Variants",
        "",
        "Fixed before the first run, all reported, none chosen afterwards. The lacunae, the vascular",
        "handling, the graph cleanup, the ownership and the ring radii are the current method's in every",
        "one, so only the named difference can move a number.",
        "",
        "**P5 is listed first because it is the preferred variant by eye. That is a visual preference,",
        "not a validated result.** No variant is called better than another by any number below.",
        "",
        "| variant | what | sigmas (um) | sigmas (px) | low cut | bridging |",
        "|---|---|---|---|---|---|",
    ]
    for key, what, family, bridging, sigmas_um, low in VARIANTS:
        if family == "current":
            lines.append(f"| {key} | {what} | | | {canaliculi.HYSTERESIS_LOW_FRACTION} | "
                         f"{'yes' if bridging else 'no'} |")
        else:
            lines.append(f"| {key} | {what} | {', '.join(str(s) for s in sigmas_um)} | "
                         f"{', '.join(str(round(params.px(s), 3)) for s in sigmas_um)} | {low} | "
                         f"{'yes' if bridging else 'no'} |")
    lines += [
        "",
        "P5 and P2 carry the same scales and the same cut: they differ by the bridging switch and",
        "nothing else, so the difference between their numbers is the effect of bridging alone.",
        "The response of `skimage.filters.sato` is already scale normalised, which was checked on",
        "synthetic ridges rather than assumed: the raw response peaks at about 0.6 times the full width",
        "at half maximum and follows the width, while multiplying by the scale squared a second time",
        "makes the largest candidate win whatever the structure.",
    ]
    lines += [
        "",
        "## The new measures",
        "",
        "- **`width_fwhm_*`**: the width measured **on the image**. At every usable skeleton pixel the",
        "  flattened intensity is sampled along the normal to the thread, and the width is the full width",
        "  at half maximum of that profile. Bridged pixels, pixels off the mask, the lacuna buffer, the",
        "  vascular mask and junction pixels are left out. **It includes the optical blur**, so it is an",
        "  apparent width in the image and an **upper bound on the diameter**, not the diameter. No point",
        "  spread function is computed or subtracted anywhere: the numerical aperture, the emission",
        "  wavelength and the pinhole are unknown.",
        "- **`width_dt_*`**: the older measure, twice the distance transform of the mask, reported beside",
        "  it. It can only ever return the width of the mask it is given.",
        "- **`contrast_snr` and `low_contrast_fraction`**: the intensity at a skeleton pixel minus the",
        "  median intensity 4 to 6 px away along the normal on both sides, divided by the noise, where the",
        "  noise is 1.4826 times the median absolute deviation of the flattened image over pixels outside",
        "  the lacunae, the vascular regions and **that variant's own mask**. The same definition is used",
        "  for every variant including A, which the earlier support measure was not: that one used the",
        "  current method's own threshold and so favoured the current method. It is kept as",
        "  `legacy_unsupported_fraction` only so the earlier tables can still be followed.",
        "- **`straight_run_fraction`** is unchanged, and **`straight_run_fraction_smoothed`** is the same",
        "  after each branch path is smoothed with a 5 point moving average, which removes most of the",
        "  effect of drawing a thin line on a pixel grid. A one pixel skeleton contains straight runs by",
        "  its nature and the current method already scores about 0.22, so only the differences between",
        "  variants carry meaning.",
        "- Everything else is measured by `src/quantification.py`, the same code for every variant.",
        "",
        "## Figures",
        "",
        "The skeleton is drawn **one pixel wide**, never thickened, on an image upscaled by nearest",
        "neighbour so a one pixel line stays visible. Owned threads take their lacuna's colour, the",
        "pipeline's own colours; everything the ownership did not reach is light grey, so **nothing is",
        "hidden**. This is the opposite of the pipeline's verification style, which thickens the skeleton",
        "and draws only what is owned.",
        "",
        "| file | what |",
        "|---|---|",
        "| `comparison_all_images.csv` and `.xlsx` | the comparison table: every measure, every image, "
        "every variant, with the summary of means and the notes |",
        "| `comparison_summary.csv` | the means over the images and the change from A |",
        "| `<label>/<label>_compare.png` | " + ", then ".join(VARIANT_KEYS)
        + f" side by side, {COMPARE_UPSCALE}x |",
        "",
        "The comparison is written as an indexed png with the skeleton colours",
        "held in reserved palette slots, so **every skeleton colour is exact** and only the background is",
        "quantised. A plain quantisation was measured first and moved all 14 overlay colours, which would",
        "have made the colours meaningless. The background error is a mean of about 1 level out of 255.",
        "",
        "## How it was run",
        "",
        "```",
        "python src/canal_physical/physical.py --dir data/WT",
        "```",
        "",
        f"Commit `{prov['git_commit']}`, tracked files clean: {not prov['git_dirty']}. Config hash "
        f"`{prov['config_hash']}`.",
        "",
        "## Every length, in micrometres and pixels",
        "",
        "| parameter | um | px | what |",
        "|---|---|---|---|",
    ]
    for name, um, px, what in params.table():
        lines.append(f"| `{name}` | {um} | {px if px is not None else ''} | {what} |")
    lines += [
        f"| `tophat_radius` (flattening) | {round(canaliculi.TOPHAT_RADIUS_PX * params.pixel_size_um, 4)} | "
        f"{canaliculi.TOPHAT_RADIUS_PX} | unchanged from the current method |",
        f"| `smooth_sigma` (flattening) | {round(canaliculi.SMOOTH_SIGMA_PX * params.pixel_size_um, 4)} | "
        f"{canaliculi.SMOOTH_SIGMA_PX} | unchanged from the current method |",
        "",
        "## Mean over the images",
        "",
        "| measure | " + " | ".join(VARIANT_KEYS) + " |",
        "|---" * (len(VARIANT_KEYS) + 1) + "|",
    ]
    for name in SUMMARY_MEASURES:
        cells = []
        for key in VARIANT_KEYS:
            value = means.get(key, {}).get(name)
            cells.append("" if value is None else f"{value:.4g}")
        lines.append(f"| `{name}` | " + " | ".join(cells) + " |")
    lines += ["", "## Versions", "", "```"]
    lines += [f"{name}: {value}" for name, value in versions.items()]
    lines += ["```", "", "## Images", ""]
    lines += [f"- {label} (`{path.name}`)" for label, path in image_paths.items()]
    path = out_dir / "README.md"
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(lines) + "\n")




# Drawing
# Moved here unchanged when src/canal_physical/figures_v2.py was removed. One
# pixel of skeleton means one pixel of the original image, blown up by whole
# blocks, so a thread is drawn at its true width and nothing is invented.

UNOWNED_COLOUR = (210, 210, 210)  # grey: a thread owned by no lacuna
SCALE_BAR_UM = 10.0
SCALE_BAR_MARGIN = 0.04
SCALE_BAR_THICKNESS = 3


def upscale(image: np.ndarray, factor: int) -> np.ndarray:
    """Nearest neighbour, so no pixel is invented and none is blurred."""
    return image.repeat(factor, axis=0).repeat(factor, axis=1)


def draw_skeleton(display: np.ndarray, skeleton: np.ndarray, owner_map: np.ndarray,
                  lacuna_id_map: np.ndarray, colors: dict, factor: int) -> np.ndarray:
    """The upscaled image with the skeleton drawn one pixel wide: owned threads in
    their lacuna's colour, everything else in light grey, lacunae outlined.

    One pixel means one pixel of the original image, so a skeleton pixel fills its
    whole block in the upscaled picture. That is the thread at its true width and
    nothing more. It is not the thickening the pipeline's verification picture
    does, which dilates the skeleton into its neighbours and makes a one pixel
    thread five pixels wide. Drawing only the centre of each block was tried first
    and is worse than useless: at eight times it puts one dot per eight pixels and
    the thread reads as a faint dotted line that disappears at any viewing size."""
    canvas = upscale(display, factor).copy()

    for lacuna_id, colour in colors.items():
        edge = segmentation.find_boundaries(lacuna_id_map == lacuna_id, mode="outer")
        if edge.any():
            canvas[upscale(edge, factor)] = colour

    unowned = upscale(skeleton & (owner_map == 0), factor)
    canvas[unowned] = UNOWNED_COLOUR
    for lacuna_id, colour in colors.items():
        owned = skeleton & (owner_map == lacuna_id)
        if owned.any():
            canvas[upscale(owned, factor)] = colour
    return canvas


def add_scale_bar(canvas: np.ndarray, pixel_size_um: float, factor: int,
                  length_um: float = SCALE_BAR_UM, colour=(255, 255, 255)) -> np.ndarray:
    """A bar of `length_um` micrometres in the bottom right, drawn into the
    pixels so nothing resamples it."""
    height, width = canvas.shape[:2]
    length = int(round(length_um / pixel_size_um * factor))
    if length >= width:
        return canvas
    margin = int(round(SCALE_BAR_MARGIN * width))
    thickness = max(1, SCALE_BAR_THICKNESS * factor // 3)
    r1 = height - margin
    c1 = width - margin
    canvas[r1 - thickness:r1, c1 - length:c1] = colour
    return canvas


def save_png(canvas: np.ndarray, path, reserved: list | None = None) -> None:
    """Write the panel, as an indexed png when `reserved` colours are given.

    A plain 256 colour quantisation moves every colour it sees, including the
    skeleton colours, which are the whole point of the picture: measured on one
    panel it changed all 14 of them. So the reserved colours keep exact slots of
    their own and only the background is quantised. The background is a dark red
    image with a narrow range, so this is hard to see, and it roughly halves the
    file."""
    from skimage.io import imsave

    path.parent.mkdir(parents=True, exist_ok=True)
    canvas = canvas.astype(np.uint8)
    if not reserved:
        imsave(path, canvas, check_contrast=False)
        return

    from PIL import Image

    unique = []
    for colour in reserved:
        colour = tuple(int(v) for v in colour)
        if colour not in unique:
            unique.append(colour)
    n_background = max(2, 256 - len(unique))
    quantised = Image.fromarray(canvas).quantize(colors=n_background, method=Image.MEDIANCUT,
                                                 dither=Image.NONE)
    palette = np.array(quantised.getpalette()[:3 * n_background], dtype=np.uint8).reshape(-1, 3)
    index = np.asarray(quantised, dtype=np.uint8).copy()
    for i, colour in enumerate(unique):
        exact = (canvas == np.array(colour, dtype=np.uint8)).all(axis=-1)
        if exact.any():
            index[exact] = n_background + i
    full = np.zeros((256, 3), dtype=np.uint8)
    full[:n_background] = palette
    for i, colour in enumerate(unique):
        full[n_background + i] = colour
    out = Image.fromarray(index, mode="P")
    out.putpalette(full.reshape(-1).tolist())
    out.save(path, optimize=True)


def pipeline_style(lac: dict, display: np.ndarray, detection: dict, path) -> None:
    """The pipeline's own verification picture, by the pipeline's own function."""
    canaliculi.save_verification({
        "lacunae": lac,
        "display": display,
        "lacuna_id_map": detection["lacuna_id_map"],
        "skeleton": detection["skeleton"],
        "owner_map": detection["owner_map"],
    }, path)


if __name__ == "__main__":
    main()
